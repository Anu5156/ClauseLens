import os
import docx
import pymupdf as fitz
from pathlib import Path
from backend.config import SAMPLE_DATA_DIR

def generate_pdf_document(file_path: Path, title: str, sections: list[tuple[str, str]]):
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)  # Standard Letter size
    margin_x = 54
    y = 54

    # Header Title
    page.insert_text((margin_x, y), title.upper(), fontsize=14, fontname="helv")
    y += 30

    for heading, text in sections:
        if y > 720:
            page = doc.new_page(width=612, height=792)
            y = 54

        if heading:
            page.insert_text((margin_x, y), heading, fontsize=11, fontname="helv")
            y += 18

        # Wrap text lines
        words = text.split(" ")
        line = ""
        for w in words:
            test_line = line + (" " if line else "") + w
            if len(test_line) > 85:
                page.insert_text((margin_x, y), line, fontsize=10, fontname="helv")
                y += 14
                if y > 720:
                    page = doc.new_page(width=612, height=792)
                    y = 54
                line = w
            else:
                line = test_line
        if line:
            page.insert_text((margin_x, y), line, fontsize=10, fontname="helv")
            y += 20

    doc.save(str(file_path))
    doc.close()

def generate_docx_document(file_path: Path, title: str, sections: list[tuple[str, str]]):
    doc = docx.Document()
    doc.add_heading(title, level=0)

    for heading, text in sections:
        if heading:
            doc.add_heading(heading, level=1)
        doc.add_paragraph(text)

    doc.save(str(file_path))

def create_all_samples():
    SAMPLE_DATA_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Residential Rental Agreement (PDF) - Flaw: Dangling cross-reference to Section 14.3
    lease_sections = [
        ("1. DEFINITIONS", 'For purposes of this Residential Lease Agreement, "Landlord" shall mean Apex Properties LLC and "Tenant" shall mean John Doe.'),
        ("2. RENT AND SECURITY DEPOSIT", '2.1 Security Deposit. The Tenant shall deposit the sum of $2,000 as a Security Deposit upon execution of this Agreement. The return of the deposit is subject to Section 14.3.'),
        ("3. OCCUPANCY AND USE", "3.1 Permitted Use. The Premises shall be used solely as a private residential dwelling for the Tenant and immediate family."),
        ("4. NOTICE AND TERMINATION", "4.1 Termination Notice. Either party may terminate this lease by providing 30 days written notice prior to the end of the term."),
        ("5. GOVERNING LAW", "5.1 Jurisdiction. This Agreement shall be governed by and construed in accordance with the laws of the State of California.")
    ]
    generate_pdf_document(
        SAMPLE_DATA_DIR / "residential_lease.pdf",
        "Residential Lease Agreement",
        lease_sections
    )

    # 2. Employment Agreement (DOCX) - Flaw: Undefined capitalized term "Restricted Territory"
    emp_sections = [
        ("1. DEFINITIONS", '1.1 Defined Terms. "Company" shall mean CloudScale Inc. "Employee" shall mean Jane Smith. "Base Salary" shall mean $150,000 per annum.'),
        ("2. DUTIES AND RESPONSIBILITIES", "2.1 Position. Employee shall serve as Senior Software Engineer reporting to the Chief Technology Officer."),
        ("3. COMPENSATION AND BENEFITS", "3.1 Payment Terms. Base Salary shall be payable in semi-monthly installments in accordance with normal payroll practices."),
        ("4. NON-COMPETE AND COVENANTS", '4.1 Non-Competition. During employment and for twelve (12) months thereafter, Employee shall not engage in any competitive business within the "Restricted Territory".'),
        ("5. TERMINATION", "5.1 Notice Period. Either party may terminate this employment by providing 30 days written notice to the other party.")
    ]
    generate_docx_document(
        SAMPLE_DATA_DIR / "employment_agreement.docx",
        "Executive Employment Agreement",
        emp_sections
    )

    # 3. SaaS Terms of Service (PDF) - Baseline Version
    saas_sections = [
        ("ARTICLE I - DEFINITIONS", 'Section 1.1 Definitions. "Service" shall mean the SaaS analytics platform. "Customer" shall mean the subscribing entity.'),
        ("ARTICLE II - SERVICE ACCESS", "Section 2.1 Grant of License. Subject to compliance with this Agreement, Company grants Customer a non-exclusive license to access the Service."),
        ("ARTICLE III - PAYMENT TERMS", "Section 3.1 Fees. Customer agrees to pay all fees specified in the Order Form within 30 days of invoice date."),
        ("ARTICLE IV - TERM AND RENEWAL", "Section 4.2 Auto-Renewal Notice. Agreement automatically renews unless Customer provides 30 days written notice prior to term expiration."),
        ("ARTICLE XI - TERMINATION FOR CONVENIENCE", "Section 11.1 Termination Notice. Customer may terminate this Agreement at any time by providing 60 days written notice to Company.")
    ]
    generate_pdf_document(
        SAMPLE_DATA_DIR / "saas_terms.pdf",
        "SaaS Terms of Service",
        saas_sections
    )

    # 4. SaaS Terms of Service Revision v2 (PDF) - Added, Removed, and Materially Changed clauses
    saas_v2_sections = [
        ("ARTICLE I - DEFINITIONS", 'Section 1.1 Definitions. "Service" shall mean the SaaS analytics platform. "Customer" shall mean the subscribing entity.'),
        ("ARTICLE II - SERVICE ACCESS", "Section 2.1 Grant of License. Subject to compliance with this Agreement, Company grants Customer a non-exclusive license to access the Service."),
        ("ARTICLE III - PAYMENT TERMS", "Section 3.1 Fees and Penalties. Customer agrees to pay all fees within 10 days of invoice date. Late payments incur a mandatory 15% flat compounding penalty per week."),
        ("ARTICLE IV - TERM AND RENEWAL", "Section 4.2 Auto-Renewal Notice. Agreement automatically renews for additional 3-year periods unless Customer provides 180 days written notice prior to expiration."),
        ("ARTICLE V - UNILATERAL MODIFICATION", "Section 5.1 Terms Modification. Company reserves the right to amend these Terms at any time by posting changes online. Customer continued use constitutes binding acceptance.")
        # Note: ARTICLE XI (Termination for Convenience) is deliberately omitted/removed in v2
    ]
    generate_pdf_document(
        SAMPLE_DATA_DIR / "saas_terms_v2.pdf",
        "SaaS Terms of Service (Revision v2)",
        saas_v2_sections
    )

if __name__ == "__main__":
    create_all_samples()
    print("Successfully generated sample documents in data/samples/")

